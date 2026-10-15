"""Formato CURSO para los videos largos: temarios, lecciones, capítulos y material de práctica.

Un curso es una serie de lecciones horizontales (16:9) agrupadas en módulos. El temario vive en
`cursos/<canal>.yaml` (fuente única: títulos, objetivos, orden y enlaces de material) y cada lección
es una ficha normal en `content/<canal>/<ID>.yaml` con `curso:` y `leccion:`.

Lo que agrega este módulo:
- datos del curso inyectados en la cabecera y en las plantillas de curso (portada, capítulo, cierre);
- capítulos de YouTube calculados con los tiempos reales del render (0:00, mínimo 3, ≥ 10 s cada uno);
- material de práctica (.xlsx con fórmulas en sintaxis española traducidas, o .txt/.md),
  verificado recalculando la planilla con LibreOffice cuando está instalado.
"""
from __future__ import annotations

import copy
import json
import re
import shutil
import subprocess
import tempfile
from functools import lru_cache
from pathlib import Path

import yaml

from .config import cargar_canal, ruta

PLANTILLAS_CURSO = {"curso_portada", "capitulo", "ejercicio", "curso_cierre"}
MARCA_ENLACE = "[ENLACE_MATERIAL]"


# ---------------------------------------------------------------- temarios

@lru_cache(maxsize=8)
def _cargar_temario(canal: str) -> dict | None:
    p = ruta("cursos", f"{canal}.yaml")
    if not p.exists():
        return None
    with open(p, encoding="utf-8") as f:
        return yaml.safe_load(f)


def temario(canal: str) -> dict | None:
    return _cargar_temario(canal)


def lecciones(canal: str) -> list[dict]:
    """Lista plana de lecciones con número, módulo y total."""
    t = temario(canal)
    if not t:
        return []
    out = []
    for m_i, mod in enumerate(t["modulos"], start=1):
        for lec in mod["lecciones"]:
            out.append({**lec, "modulo_n": m_i, "modulo": mod["titulo"]})
    for n, lec in enumerate(out, start=1):
        lec["n"] = n
        lec["total"] = len(out)
    return out


def leccion_de(spec: dict) -> dict | None:
    if not spec.get("curso"):
        return None
    t = temario(spec["canal"])
    if not t or t.get("id") != spec["curso"]:
        return None
    ls = lecciones(spec["canal"])
    for lec in ls:
        if lec["id"] == spec["id"]:
            lec = dict(lec)
            lec["curso_titulo"] = t["titulo"]
            lec["curso_id"] = t["id"]
            sig = next((x for x in ls if x["n"] == lec["n"] + 1), None)
            lec["siguiente"] = sig
            return lec
    return None


# ---------------------------------------------------------------- capítulos

def titulo_capitulo(escena: dict) -> str | None:
    if escena.get("capitulo"):
        return str(escena["capitulo"])
    v = escena.get("visual") or {}
    if v.get("plantilla") == "capitulo":
        return str((v.get("datos") or {}).get("titulo", "")) or None
    return None


def indices_capitulos(spec: dict) -> list[tuple[int, str]]:
    """(índice de escena, título). Si la primera escena no abre capítulo, se agrega 'Introducción'."""
    caps = [(i, t) for i, e in enumerate(spec["escenas"]) if (t := titulo_capitulo(e))]
    if not caps or caps[0][0] != 0:
        caps.insert(0, (0, "Introducción"))
    return caps


def _sin_marcas(t: str) -> str:
    t = re.sub(r"\[\[([^|\]]+)\|([^\]]+)\]\]", r"\2", t)
    return t.replace("*", "").strip()


def fmt_tiempo(s: float, con_horas: bool = False) -> str:
    s = int(s)
    h, m, sec = s // 3600, (s % 3600) // 60, s % 60
    return f"{h}:{m:02d}:{sec:02d}" if con_horas else f"{m}:{sec:02d}"


def capitulos(spec: dict, tecnico: dict | None) -> list[dict]:
    """Capítulos con tiempos reales del render. Vacío si el video aún no se produjo."""
    if not tecnico or not tecnico.get("escenas"):
        return []
    ini = [e["inicio"] for e in tecnico["escenas"]]
    dur = float(tecnico.get("duracion_s") or 0)
    caps = indices_capitulos(spec)
    out = []
    for k, (i, t) in enumerate(caps):
        if i >= len(ini):
            continue
        a = 0.0 if k == 0 else ini[i]
        b = ini[caps[k + 1][0]] if k + 1 < len(caps) and caps[k + 1][0] < len(ini) else dur
        out.append({"inicio": a, "fin": b, "titulo": _sin_marcas(t)})
    return out


def validar_capitulos(caps: list[dict]) -> list[str]:
    err = []
    if len(caps) < 3:
        err.append(f"YouTube exige al menos 3 capítulos (hay {len(caps)})")
    if caps and caps[0]["inicio"] != 0:
        err.append("el primer capítulo debe empezar en 0:00")
    for c in caps:
        if c["fin"] - c["inicio"] < 10:
            err.append(f"capítulo «{c['titulo']}» dura menos de 10 s")
    return err


def texto_capitulos(caps: list[dict]) -> str:
    horas = bool(caps) and caps[-1]["fin"] >= 3600
    return "\n".join(f"{fmt_tiempo(c['inicio'], horas)} {c['titulo']}" for c in caps)


def leer_tecnico(spec: dict) -> dict | None:
    p = cargar_canal(spec["canal"]).dir_salida("metadata") / f"{spec['id']}.tecnico.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


# ---------------------------------------------------------------- preparación para el render

def preparar(spec: dict) -> dict:
    """Copia de la ficha con los datos del curso inyectados en cabecera y plantillas de curso."""
    lec = leccion_de(spec)
    if not lec:
        return spec
    s = copy.deepcopy(spec)
    caps = indices_capitulos(s)
    inicio_cap = {i: k for k, (i, _) in enumerate(caps)}
    actual = 0
    material = [m["archivo"] for m in s.get("materiales", [])]
    for i, e in enumerate(s["escenas"]):
        if i in inicio_cap:
            actual = inicio_cap[i]
        v = e["visual"]
        v["_cab"] = {"curso": lec["curso_titulo"], "leccion": lec["n"], "total": lec["total"],
                     "modulo_n": lec["modulo_n"], "modulo": lec["modulo"],
                     "cap_n": actual + 1, "cap_total": len(caps)}
        d = v.setdefault("datos", {})
        if v["plantilla"] == "curso_portada":
            d.setdefault("curso", lec["curso_titulo"])
            d.setdefault("leccion", lec["n"])
            d.setdefault("total", lec["total"])
            d.setdefault("modulo", f"Módulo {lec['modulo_n']}: {lec['modulo']}")
            d.setdefault("titulo", lec["titulo"])
            d.setdefault("objetivos", lec.get("objetivos", []))
        elif v["plantilla"] == "capitulo":
            d.setdefault("n", actual + 1)
            d.setdefault("total", len(caps))
        elif v["plantilla"] == "curso_cierre":
            sig = lec.get("siguiente")
            if sig:
                d.setdefault("siguiente", f"Lección {sig['n']}: {sig['titulo']}")
            else:
                d.setdefault("siguiente", "Fin del curso")
            if material:
                d.setdefault("material", material[0])
        elif v["plantilla"] == "ejercicio" and material:
            d.setdefault("material", material[0])
    return s


# ---------------------------------------------------------------- material de práctica

FUNCIONES_ES_EN = {
    "SUMA": "SUM", "PROMEDIO": "AVERAGE", "CONTAR": "COUNT", "CONTARA": "COUNTA", "CONTAR.BLANCO": "COUNTBLANK",
    "MAX": "MAX", "MIN": "MIN", "SI": "IF", "Y": "AND", "O": "OR", "NO": "NOT", "SI.ERROR": "IFERROR",
    "CONTAR.SI": "COUNTIF", "SUMAR.SI": "SUMIF", "PROMEDIO.SI": "AVERAGEIF",
    "CONTAR.SI.CONJUNTO": "COUNTIFS", "SUMAR.SI.CONJUNTO": "SUMIFS", "PROMEDIO.SI.CONJUNTO": "AVERAGEIFS",
    "BUSCARV": "VLOOKUP", "BUSCARH": "HLOOKUP", "INDICE": "INDEX", "COINCIDIR": "MATCH",
    "REDONDEAR": "ROUND", "REDONDEAR.MAS": "ROUNDUP", "REDONDEAR.MENOS": "ROUNDDOWN", "ENTERO": "INT", "ABS": "ABS",
    "HOY": "TODAY", "AHORA": "NOW", "FECHA": "DATE", "AÑO": "YEAR", "MES": "MONTH", "DIA": "DAY", "DÍA": "DAY",
    "DIAS.LAB": "NETWORKDAYS", "SIFECHA": "DATEDIF", "FIN.MES": "EOMONTH",
    "PAGO": "PMT", "VF": "FV", "VA": "PV", "NPER": "NPER", "TASA": "RATE",
    "CONCATENAR": "CONCATENATE", "IZQUIERDA": "LEFT", "DERECHA": "RIGHT", "EXTRAE": "MID", "LARGO": "LEN",
    "MAYUSC": "UPPER", "MINUSC": "LOWER", "NOMPROPIO": "PROPER", "ESPACIOS": "TRIM", "TEXTO": "TEXT", "VALOR": "VALUE",
    "SUSTITUIR": "SUBSTITUTE", "HALLAR": "SEARCH", "ENCONTRAR": "FIND", "ESBLANCO": "ISBLANK", "ESNUMERO": "ISNUMBER",
    "PRODUCTO": "PRODUCT", "SUMAPRODUCTO": "SUMPRODUCT", "RAIZ": "SQRT", "POTENCIA": "POWER", "RESIDUO": "MOD",
    "K.ESIMO.MAYOR": "LARGE", "K.ESIMO.MENOR": "SMALL", "JERARQUIA": "RANK", "MEDIANA": "MEDIAN",
    # funciones nuevas: Excel exige el prefijo _xlfn en el archivo
    "BUSCARX": "_xlfn.XLOOKUP", "COINCIDIRX": "_xlfn.XMATCH", "FILTRAR": "_xlfn._xlws.FILTER",
    "UNICOS": "_xlfn.UNIQUE", "ORDENAR": "_xlfn._xlws.SORT", "SECUENCIA": "_xlfn.SEQUENCE",
    "SI.CONJUNTO": "_xlfn.IFS", "CONCAT": "_xlfn.CONCAT", "UNIRCADENAS": "_xlfn.TEXTJOIN", "LET": "_xlfn.LET",
}
# LibreOffice 24.2 no calcula estas (24.8+ sí): se marcan como "sin verificar" en lugar de error
NO_VERIFICABLES_LO = {"_xlfn.XLOOKUP", "_xlfn.XMATCH", "_xlfn._xlws.FILTER", "_xlfn.UNIQUE", "_xlfn._xlws.SORT",
                      "_xlfn.SEQUENCE", "_xlfn.LET"}


class ErrorFormula(ValueError):
    pass


def formula_a_excel(f: str) -> str:
    """Traduce una fórmula escrita como en Excel en español (=SUMA(B2:B5;C1)) a la forma interna del
    archivo .xlsx (=SUM(B2:B5,C1)). Excel la muestra traducida al idioma de cada usuario."""
    if not f.startswith("="):
        return f
    partes = re.split(r'("(?:[^"]|"")*")', f)  # no tocar textos entre comillas
    out = []
    for k, p in enumerate(partes):
        if k % 2 == 1:
            out.append(p)
            continue
        p = re.sub(r"(?<=\d),(?=\d)", ".", p)  # coma decimal -> punto
        p = p.replace(";", ",")

        def fn(m):
            nombre = m.group(1).upper()
            if nombre not in FUNCIONES_ES_EN:
                raise ErrorFormula(f"función no reconocida: {m.group(1)} (agregarla a FUNCIONES_ES_EN)")
            return FUNCIONES_ES_EN[nombre] + "("
        p = re.sub(r"(?<![A-Za-z_.])([A-ZÁÉÍÓÚÑa-záéíóúñ][A-ZÁÉÍÓÚÑa-záéíóúñ.]*)\(", fn, p)
        p = re.sub(r"\bVERDADERO\b(?!\()", "TRUE", p)
        p = re.sub(r"\bFALSO\b(?!\()", "FALSE", p)
        out.append(p)
    return "".join(out)


def _hex(c: str) -> str:
    return c.lstrip("#").upper()


def crear_xlsx(mat: dict, canal, destino: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    p = canal.paleta
    wb = Workbook()
    wb.remove(wb.active)
    fino = Side(style="thin", color="D9DCD5")
    for hoja in mat["hojas"]:
        ws = wb.create_sheet(hoja["nombre"][:31])
        fila0 = 1
        if hoja.get("instrucciones"):
            ws.cell(1, 1, hoja["instrucciones"]).font = Font(italic=True, color=_hex(p["texto_suave"]))
            fila0 = 3
        for r, fila in enumerate(hoja["filas"]):
            for c, v in enumerate(fila):
                if v is None or v == "":
                    continue
                celda = ws.cell(fila0 + r, c + 1, formula_a_excel(v) if isinstance(v, str) else v)
                celda.border = Border(top=fino, bottom=fino, left=fino, right=fino)
                if r == 0 and hoja.get("encabezado", True):
                    celda.font = Font(bold=True, color=_hex(p.get("texto_sobre_acento", "#ffffff")))
                    celda.fill = PatternFill("solid", fgColor=_hex(p["acento"]))
                    celda.alignment = Alignment(horizontal="center")
        for col, fmt in (hoja.get("formatos") or {}).items():
            for r in range(fila0 + 1, fila0 + len(hoja["filas"])):
                ws[f"{col}{r}"].number_format = fmt
        for i, ancho in enumerate(hoja.get("anchos") or [], start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
        if hoja.get("encabezado", True):
            ws.freeze_panes = ws.cell(fila0 + 1, 1)
    wb.properties.creator = canal.nombre
    wb.properties.title = mat.get("titulo", mat["archivo"])
    wb.save(destino)


def _celda_destino(hoja: dict, ref: str) -> str:
    """Las referencias de 'verificar' son relativas a la tabla; si hay instrucciones arriba, se desplazan 2 filas."""
    if not hoja.get("instrucciones"):
        return ref
    m = re.match(r"([A-Z]+)(\d+)$", ref)
    return f"{m.group(1)}{int(m.group(2)) + 2}"


def verificar_xlsx(mat: dict, archivo: Path) -> dict:
    """Recalcula la planilla con LibreOffice (si está instalado) y compara con los valores esperados."""
    esperados = [(h, ref, val) for h in mat["hojas"] for ref, val in (h.get("verificar") or {}).items()]
    if not esperados:
        return {"estado": "sin_controles", "detalle": "la ficha no declara valores a verificar"}
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return {"estado": "sin_verificar", "detalle": "LibreOffice no está instalado (opcional)"}
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as tmp:
        perfil = Path(tmp, "perfil").as_uri()
        r = subprocess.run([soffice, f"-env:UserInstallation={perfil}", "--headless", "--calc", "--convert-to", "xlsx",
                            "--outdir", tmp, str(archivo)], capture_output=True, text=True, timeout=180)
        calc = Path(tmp, archivo.name)
        if r.returncode != 0 or not calc.exists():
            return {"estado": "sin_verificar", "detalle": f"LibreOffice no pudo recalcular: {r.stderr[-200:]}"}
        wb_v = load_workbook(calc, data_only=True)
        wb_f = load_workbook(archivo)
        fallos, sin_verif, ok = [], [], 0
        for h, ref, esperado in esperados:
            celda = _celda_destino(h, ref)
            nombre = h["nombre"][:31]
            valor = wb_v[nombre][celda].value
            formula = str(wb_f[nombre][celda].value or "")
            if isinstance(valor, str) and valor.startswith("#") and any(fn in formula for fn in NO_VERIFICABLES_LO):
                sin_verif.append(f"{nombre}!{ref}")
                continue
            if isinstance(esperado, (int, float)) and isinstance(valor, (int, float)):
                bien = abs(valor - esperado) <= max(0.006, abs(esperado) * 1e-6)
            else:
                bien = str(valor) == str(esperado)
            if bien:
                ok += 1
            else:
                fallos.append(f"{nombre}!{ref}: {valor!r} (esperado {esperado!r}; fórmula {formula})")
    estado = "error" if fallos else ("parcial" if sin_verif else "ok")
    return {"estado": estado, "ok": ok, "fallos": fallos, "sin_verificar": sin_verif,
            "detalle": "; ".join(fallos) or (f"{ok} valores comprobados" + (f"; sin verificar con este LibreOffice: {', '.join(sin_verif)}" if sin_verif else ""))}


def generar_materiales(spec: dict) -> list[dict]:
    """Crea los archivos de práctica de la lección en output/<canal>/materiales/<ID>/ y los verifica."""
    mats = spec.get("materiales") or []
    if not mats:
        return []
    canal = cargar_canal(spec["canal"])
    carpeta = canal.dir_salida("materiales", spec["id"])
    out = []
    for m in mats:
        destino = carpeta / m["archivo"]
        if m["tipo"] == "xlsx":
            crear_xlsx(m, canal, destino)
            ver = verificar_xlsx(m, destino)
        elif m["tipo"] in ("texto", "md", "txt"):
            destino.write_text(m["contenido"].strip() + "\n", encoding="utf-8")
            ver = {"estado": "ok", "detalle": "archivo de texto"}
        else:
            raise ValueError(f"tipo de material desconocido: {m['tipo']}")
        out.append({"archivo": m["archivo"], "ruta": str(destino), "bytes": destino.stat().st_size, **ver})
    return out


def enlace_material(spec: dict) -> str | None:
    """Enlace público del material (lo carga el usuario en cursos/<canal>.yaml después de subirlo a su Drive)."""
    lec = leccion_de(spec)
    if not lec or not spec.get("materiales"):
        return None
    return (lec.get("enlace_material") or "").strip() or MARCA_ENLACE


# ---------------------------------------------------------------- validación de la ficha de lección

def validar_leccion(spec: dict) -> list[str]:
    err = []
    lec = leccion_de(spec)
    if not lec:
        return [f"la lección {spec['id']} no figura en el temario cursos/{spec['canal']}.yaml (curso '{spec.get('curso')}')"]
    if spec["formato"] != "largo":
        err.append("una lección de curso debe tener formato 'largo'")
    if spec["escenas"][0]["visual"]["plantilla"] != "curso_portada":
        err.append("la lección debe empezar con la plantilla 'curso_portada'")
    if spec["escenas"][-1]["visual"]["plantilla"] != "curso_cierre":
        err.append("la lección debe terminar con la plantilla 'curso_cierre'")
    if not any(e["visual"]["plantilla"] == "ejercicio" for e in spec["escenas"]):
        err.append("la lección necesita al menos una escena 'ejercicio' (práctica)")
    if len(indices_capitulos(spec)) < 3:
        err.append("la lección necesita al menos 3 capítulos (YouTube los exige para mostrarlos)")
    for m in spec.get("materiales") or []:
        if m.get("tipo") == "xlsx":
            for h in m.get("hojas", []):
                for fila in h.get("filas", []):
                    for v in fila:
                        if isinstance(v, str) and v.startswith("="):
                            try:
                                formula_a_excel(v)
                            except ErrorFormula as e:
                                err.append(f"material {m['archivo']}: {e}")
    return err
